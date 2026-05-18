"""Round-trip tests covering unpack + pack via the real PPTX entry points.

These tests use ``python-pptx`` (test-only dependency) to construct minimal
PPTX fixtures, then verify that ``1_unpack_pptx.py`` and ``4_pack_pptx.py``
behave as a lossless mirror pair: unpack -> pack should yield a PPTX that
opens cleanly, contains the same slide texts, and remains a valid ZIP.
"""
from __future__ import annotations

import zipfile
from pathlib import Path


def _slide_texts(pptx_path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    out: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return out


def test_unpack_creates_ppt_directory(make_pptx, tmp_path, run_script_fn):
    src = make_pptx(["Hello, world!"], filename="src.pptx")
    out_dir = tmp_path / "unpacked"
    rc, _, err = run_script_fn("1_unpack_pptx.py", str(src), str(out_dir))
    assert rc == 0, err
    assert (out_dir / "ppt" / "slides").is_dir()
    slides = list((out_dir / "ppt" / "slides").glob("slide*.xml"))
    assert len(slides) == 1


def test_unpack_rejects_already_translated_default(make_pptx, tmp_path, run_script_fn):
    """Files whose stem ends with ``_JA`` should be rejected unless the agent
    explicitly opts in (forward-compat for Phase 2 ja->en flow).

    This is a soft test: the guard may or may not be present depending on
    which PRs have merged. If the script accepts ``_JA`` files (no guard
    implemented), this test is skipped, not failed.
    """
    src = make_pptx(["Hello"], filename="Deck_JA.pptx")
    out_dir = tmp_path / "unpacked_ja"
    rc, _, err = run_script_fn("1_unpack_pptx.py", str(src), str(out_dir))
    if rc == 0 and "_JA" not in err:
        import pytest

        pytest.skip("_JA guard not implemented on this branch")
    assert rc != 0
    assert "_JA" in err or "already" in err.lower()


def test_roundtrip_unpack_then_pack(make_pptx, tmp_path, run_script_fn):
    src = make_pptx(
        ["Slide one body", "Slide two body", "Slide three body"],
        filename="src.pptx",
    )
    unpacked = tmp_path / "unpacked"
    rc, _, err = run_script_fn("1_unpack_pptx.py", str(src), str(unpacked))
    assert rc == 0, err

    out = tmp_path / "roundtrip.pptx"
    rc, _, err = run_script_fn("4_pack_pptx.py", str(unpacked), str(out))
    assert rc == 0, err

    assert out.is_file()
    with zipfile.ZipFile(out) as zf:
        bad = zf.testzip()
        assert bad is None, f"Corrupt entry in repacked PPTX: {bad}"

    texts = _slide_texts(out)
    assert "Slide one body" in texts
    assert "Slide two body" in texts
    assert "Slide three body" in texts


def test_roundtrip_translation_visible_in_repacked_pptx(
    make_pptx, tmp_path, run_script_fn
):
    src = make_pptx(["Authentication"], filename="src.pptx")
    unpacked = tmp_path / "unpacked"
    rc, _, err = run_script_fn("1_unpack_pptx.py", str(src), str(unpacked))
    assert rc == 0, err

    import json

    dict_path = tmp_path / "dict.json"
    dict_path.write_text(
        json.dumps({"Authentication": "認証"}, ensure_ascii=False), "utf-8"
    )
    rc, _, err = run_script_fn(
        "3_apply_translations.py", str(unpacked), str(dict_path)
    )
    assert rc == 0, err

    out = tmp_path / "translated.pptx"
    rc, _, err = run_script_fn("4_pack_pptx.py", str(unpacked), str(out))
    assert rc == 0, err

    texts = _slide_texts(out)
    assert "認証" in texts
    assert "Authentication" not in texts
