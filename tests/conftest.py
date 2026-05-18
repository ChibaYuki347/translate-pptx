"""Pytest fixtures for translate-pptx tests.

Generates minimal PPTX fixtures programmatically with python-pptx so the test
suite is self-contained (no binary fixtures checked into the repo).
"""
from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Iterable

import pytest


REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / ".github" / "agents" / "scripts"

# Make the production scripts importable as plain modules in tests.
sys.path.insert(0, str(SCRIPTS_DIR))


@pytest.fixture(scope="session")
def scripts_dir() -> Path:
    """Return the absolute path to ``.github/agents/scripts``."""
    return SCRIPTS_DIR


@pytest.fixture
def make_pptx(tmp_path: Path):
    """Factory fixture returning a function that builds a minimal PPTX file.

    Usage::

        path = make_pptx([
            "Hello, world!",
            "Authentication is required.",
        ])
    """
    from pptx import Presentation

    def _make(texts: Iterable[str], *, filename: str = "fixture.pptx") -> Path:
        prs = Presentation()
        blank = prs.slide_layouts[6]  # blank layout
        for t in texts:
            slide = prs.slides.add_slide(blank)
            tx = slide.shapes.add_textbox(914400, 914400, 7315200, 1828800)
            tx.text_frame.text = t
        out = tmp_path / filename
        prs.save(out)
        return out

    return _make


@pytest.fixture
def unpacked_dir(tmp_path: Path) -> Path:
    """Create a minimal unpacked PPTX directory tree (no slides yet)."""
    root = tmp_path / "unpacked"
    (root / "ppt" / "slides").mkdir(parents=True)
    (root / "ppt" / "notesSlides").mkdir(parents=True)
    (root / "ppt" / "slideLayouts").mkdir(parents=True)
    (root / "ppt" / "slideMasters").mkdir(parents=True)
    (root / "ppt" / "notesMasters").mkdir(parents=True)
    (root / "ppt" / "theme").mkdir(parents=True)
    return root


def write_slide(
    unpacked: Path,
    slide_index: int,
    runs: Iterable[tuple[str, str]],
    *,
    in_notes: bool = False,
) -> Path:
    """Write a minimal slideN.xml or notesSlideN.xml with the given runs.

    Each *runs* entry is ``(lang_attribute_value, text)``. ``lang_attribute_value``
    is inserted as ``lang="..."``; use the empty string to omit it.
    """
    sub = "notesSlides" if in_notes else "slides"
    name = ("notesSlide" if in_notes else "slide") + str(slide_index) + ".xml"
    body = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>']
    body.append(
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    )
    body.append("<p:cSld><p:spTree>")
    for lang, text in runs:
        body.append("<p:sp><p:txBody><a:p><a:r>")
        if lang:
            body.append(f'<a:rPr lang="{lang}"><a:latin typeface="Arial"/></a:rPr>')
        else:
            body.append('<a:rPr><a:latin typeface="Arial"/></a:rPr>')
        body.append(f"<a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>")
    body.append("</p:spTree></p:cSld></p:sld>")
    path = unpacked / "ppt" / sub / name
    path.write_text("\n".join(body), encoding="utf-8")
    return path


@pytest.fixture
def write_slide_fn():
    """Expose ``write_slide`` as a pytest fixture for readability in tests."""
    return write_slide


def run_script(script_name: str, *args: str) -> tuple[int, str, str]:
    """Run one of the pipeline scripts as a subprocess and return (rc, stdout, stderr)."""
    import subprocess

    script = SCRIPTS_DIR / script_name
    proc = subprocess.run(
        [sys.executable, str(script), *args],
        capture_output=True,
        text=True,
        check=False,
    )
    return proc.returncode, proc.stdout, proc.stderr


@pytest.fixture
def run_script_fn():
    """Expose ``run_script`` for tests that need to invoke scripts as subprocesses."""
    return run_script
